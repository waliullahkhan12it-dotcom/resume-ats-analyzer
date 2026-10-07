import json
import os
import tempfile
from pathlib import Path

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader


MODEL_NAME = "gemini-2.5-flash"


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(__import__("io").BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        doc = Document(tmp_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs).strip()
    finally:
        Path(tmp_path).unlink(missing_ok=True)


def extract_resume_text(uploaded_file) -> str:
    data = uploaded_file.getvalue()
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(data)
    if suffix == ".docx":
        return extract_docx_text(data)
    raise ValueError("Unsupported file type. Please upload a PDF or DOCX file.")


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict:
    client = genai.Client(api_key=api_key)

    job_context = (
        job_description.strip()
        if job_description.strip()
        else "No job description was provided. Evaluate general ATS readiness and resume quality."
    )

    prompt = f"""
You are an expert ATS resume reviewer and career coach.

Analyze the resume below. Produce a practical ATS-readiness assessment.

IMPORTANT:
- The score is an ESTIMATE, not a score from a real ATS vendor.
- If a job description is provided, compare the resume against it.
- Do not invent experience, education, skills, dates, employers, or achievements.
- Penalize missing/weak sections, unclear formatting, lack of measurable achievements,
  weak keyword alignment, and ATS-unfriendly structures.
- Reward relevant keywords, clear section headings, measurable results, concise bullets,
  standard terminology, and readable structure.
- Give actionable improvements.

Return ONLY valid JSON matching the requested schema.

JOB DESCRIPTION:
{job_context}

RESUME:
{resume_text}
"""

    schema = {
        "type": "object",
        "properties": {
            "ats_score": {
                "type": "integer",
                "description": "Estimated ATS readiness score from 0 to 100."
            },
            "summary": {
                "type": "string",
                "description": "Short overall assessment."
            },
            "section_scores": {
                "type": "object",
                "properties": {
                    "keywords": {"type": "integer"},
                    "formatting": {"type": "integer"},
                    "experience": {"type": "integer"},
                    "skills": {"type": "integer"},
                    "education": {"type": "integer"},
                    "impact": {"type": "integer"}
                },
                "required": [
                    "keywords", "formatting", "experience",
                    "skills", "education", "impact"
                ]
            },
            "strengths": {
                "type": "array",
                "items": {"type": "string"}
            },
            "improvements": {
                "type": "array",
                "items": {"type": "string"}
            },
            "missing_keywords": {
                "type": "array",
                "items": {"type": "string"}
            },
            "ats_warnings": {
                "type": "array",
                "items": {"type": "string"}
            },
            "rewritten_examples": {
                "type": "array",
                "items": {"type": "string"}
            }
        },
        "required": [
            "ats_score", "summary", "section_scores", "strengths",
            "improvements", "missing_keywords", "ats_warnings",
            "rewritten_examples"
        ]
    }

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=schema,
            temperature=0.2,
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    result = json.loads(response.text)
    result["ats_score"] = max(0, min(100, int(result["ats_score"])))
    return result


def score_label(score: int) -> str:
    if score >= 85:
        return "Excellent ATS readiness"
    if score >= 70:
        return "Good ATS readiness"
    if score >= 50:
        return "Needs improvement"
    return "Major improvements recommended"


st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

st.title("📄 Resume ATS Analyzer")
st.caption("Upload a resume and get an AI-powered ATS-readiness estimate and improvement plan.")

with st.sidebar:
    st.header("Settings")
    api_key = st.text_input(
        "Gemini API key",
        type="password",
        value=os.getenv("GEMINI_API_KEY", ""),
        help="For deployment, store this as a Streamlit secret named GEMINI_API_KEY.",
    )
    st.info(
        "The ATS score is an AI-based estimate. Different ATS systems use different "
        "rules, so it should be used as guidance rather than a guaranteed hiring score."
    )

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="PDF or DOCX only.",
)

job_description = st.text_area(
    "Optional: paste the job description",
    height=220,
    placeholder="Adding the target job description makes keyword matching much more useful.",
)

if uploaded_file:
    st.success(f"Uploaded: {uploaded_file.name}")

    if st.button("🔍 Analyze Resume", type="primary", use_container_width=True):
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            st.stop()

        try:
            with st.spinner("Extracting resume and analyzing it with Gemini..."):
                resume_text = extract_resume_text(uploaded_file)

                if len(resume_text.strip()) < 80:
                    st.error(
                        "Very little text could be extracted. If this is a scanned/image PDF, "
                        "please use a text-based PDF or DOCX."
                    )
                    st.stop()

                result = analyze_resume(resume_text, job_description, api_key)

            st.session_state["analysis"] = result
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")

analysis = st.session_state.get("analysis")

if analysis:
    st.divider()
    st.subheader("ATS Score")

    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Estimated ATS Score", f"{analysis['ats_score']}/100")
    with col2:
        st.progress(analysis["ats_score"] / 100)
        st.write(score_label(analysis["ats_score"]))

    st.write(analysis["summary"])

    st.subheader("Section Scores")
    scores = analysis["section_scores"]
    cols = st.columns(len(scores))
    for col, (name, value) in zip(cols, scores.items()):
        with col:
            st.metric(name.replace("_", " ").title(), f"{value}/100")

    left, right = st.columns(2)

    with left:
        st.subheader("✅ Strengths")
        for item in analysis["strengths"]:
            st.markdown(f"- {item}")

        st.subheader("🔑 Missing / Weak Keywords")
        if analysis["missing_keywords"]:
            for item in analysis["missing_keywords"]:
                st.markdown(f"- `{item}`")
        else:
            st.write("No major keyword gaps were identified.")

    with right:
        st.subheader("⚠️ ATS Warnings")
        if analysis["ats_warnings"]:
            for item in analysis["ats_warnings"]:
                st.markdown(f"- {item}")
        else:
            st.write("No major ATS warnings were identified.")

        st.subheader("🚀 Improvements")
        for item in analysis["improvements"]:
            st.markdown(f"- {item}")

    st.subheader("✍️ Rewrite Examples")
    for item in analysis["rewritten_examples"]:
        st.markdown(f"- {item}")

    st.download_button(
        "Download analysis as JSON",
        data=json.dumps(analysis, indent=2),
        file_name="resume_ats_analysis.json",
        mime="application/json",
    )
